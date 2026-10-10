from pathlib import Path
import os,sys,json,subprocess,hashlib,gzip,shutil,tempfile,time,random,statistics,zipfile,xml.etree.ElementTree as ET,re,platform,tarfile
ROOT=Path.cwd();TASK=ROOT/'validation-macos-pgo';TARGET=os.environ['MACOS_PGO_TARGET'];OUT=TASK/'evidence'/TARGET;OUT.mkdir(parents=True,exist_ok=True)
SOURCE_COMMIT='cdd9b13c41502efaed4bfc6a6b868af1a7bc0ec0';VERSION='1.99.0';LOOPS=4;PAIRS=64
report={'source_commit':SOURCE_COMMIT,'workflow_commit':os.environ.get('GITHUB_SHA'),'run':os.environ.get('GITHUB_RUN_ID'),'target':TARGET,'criteria':{'primary_paired_median_at_most':-.01,'primary_upper95_below':0,'workbook_non_regression_upper95':.03,'pairs':PAIRS,'invocations_per_sample':LOOPS,'bootstrap':10000,'seed':2026101014,'size':'record actual executable/package; no zero-growth requirement'},'cases':{},'checks':[]}
def save(): (OUT/'results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def cmd(args,cwd=ROOT,env=None,expected=0,timeout=300):
 p=subprocess.run(list(map(str,args)),cwd=cwd,env=env,capture_output=True,timeout=timeout)
 if expected is not None:assert p.returncode==expected,(list(map(str,args)),p.returncode,p.stdout[-1000:],p.stderr[-7000:])
 return p
def cargo(args,td,flags=None):
 e=dict(os.environ,CARGO_TARGET_DIR=str(td));e.pop('RUSTFLAGS',None);e.pop('CARGO_ENCODED_RUSTFLAGS',None)
 if flags:e['RUSTFLAGS']=flags
 p=cmd(['cargo','+'+VERSION,'--verbose',*args],env=e)
 text=p.stderr.decode('utf-8','replace')
 with (OUT/'build.log').open('a') as f:f.write('cargo '+str(args)+'\n'+text+'\n')
 assert 'hash mismatch' not in text.lower(),text
 missing=[line for line in text.splitlines() if 'no profile data available for function' in line]
 for line in missing:
  known=('3std2rt10lang_startI' in line or ('9drop_glue' in line and '3VechEE' in line) or ('zip_archive' in line and '4reads4_0' in line) or ('xlsx_container' in line and '19from_validated_file0' in line) or ('6RawVec' in line and '8grow_one' in line and any(x in line for x in ['DeflateToken','RawVecReE','RawVecTRe'])))
  assert known and 'up to 0 count discarded' in line,('Unexpected missing-profile diagnostic',line)
 if missing:
  report.setdefault('profile_diagnostics',[]).append({'command':args,'missing_functions':missing,'hash_mismatches':0,'review':'Known missing entries (std generic/closure helpers) with zero counts discarded; default compiler warning remains enabled; actual workload validation required.'});save()
 return p
report['environment']={'rust':cmd(['rustc','+'+VERSION,'-Vv']).stdout.decode(),'os':platform.platform(),'cpu':cmd(['sysctl','-n','machdep.cpu.brand_string'],expected=None).stdout.decode().strip(),'arch':platform.machine(),'runner_image':os.environ.get('ImageVersion')}
assert cmd(['rustc','+'+VERSION,'-vV']).stdout.decode().split('release: ')[1].splitlines()[0]==VERSION
host=next(x.split(': ',1)[1] for x in report['environment']['rust'].splitlines() if x.startswith('host:'));assert host==TARGET,(host,TARGET)
report['sources']={str(p.relative_to(ROOT)):digest(p) for p in sorted((ROOT/'src').rglob('*')) if p.is_file()}
report['lock_sha256']=digest(ROOT/'Cargo.lock');report['linux_windows_profiles']={p.name:digest(p) for p in (ROOT/'pgo').glob('*.profdata')}
report['fixtures']={name:digest(TASK/name) for name in ['before.xlsx','valid.xlsx','source.xls.gz']};report['fixtures']['current.xlsx']=digest(ROOT/'fuel_cost_chungcheong.xlsx')
save()
# Actual unchanged CLI, default features, identical explicit native targets and release profile.
normal_td=TASK/'build-normal';train_td=TASK/'build-train';pgo_td=TASK/'build-pgo';raw=TASK/'raw';raw.mkdir(exist_ok=True)
cargo(['fmt','--all','--','--check'],normal_td)
cargo(['clippy','--release','--frozen','--all-targets','--target',TARGET,'--','-D','warnings'],normal_td)
cargo(['build','--release','--frozen','--bin','fcupdater','--target',TARGET],normal_td)
normal=normal_td/TARGET/'release/fcupdater'
source=TASK/'source.xls';source.write_bytes(gzip.decompress((TASK/'source.xls.gz').read_bytes()))
replay=TASK/'replay.dylib';cmd(['clang','-dynamiclib','-O2',TASK/'replay.c','-lcurl','-o',replay])
env=dict(os.environ,PGO_XLS=str(source),DYLD_INSERT_LIBRARIES=str(replay),DYLD_FORCE_FLAT_NAMESPACE='1');env.pop('LLVM_PROFILE_FILE',None)
trainenv=dict(env,LLVM_PROFILE_FILE=str(raw/'%p-%m.profraw'))
cargo(['build-pgo'],train_td,'-Cprofile-generate='+str(raw))
train=train_td/TARGET/'release/fcupdater'
def canonical(path):
 with zipfile.ZipFile(path) as z:
  assert z.testzip() is None
  result={}
  for name in sorted(z.namelist()):
   b=z.read(name)
   if name=='docProps/core.xml':
    root=ET.fromstring(b)
    for e in root.iter():
     if e.tag.endswith('}modified'):e.text='NORMALIZED'
    b=ET.tostring(root)
   result[name]=digest_bytes(b)
  return result
def digest_bytes(b):return hashlib.sha256(b).hexdigest()
def update(exe,fixture,verify,e=env,loops=1,independent=False):
 with tempfile.TemporaryDirectory(prefix='fcupdater-macos-pgo-') as d:
  work=Path(d);target=work/'fuel_cost_chungcheong.xlsx';elapsed=0;sig=None
  for _ in range(loops):
   shutil.copy2(fixture,target);start=time.perf_counter();p=cmd([exe,*(['--verify'] if verify else [])],work,e);elapsed+=time.perf_counter()-start
   sig=(p.stdout.hex(),p.stderr.hex(),canonical(target))
  if independent:
   from openpyxl import load_workbook
   book=load_workbook(target,read_only=True,data_only=False)
   assert book.sheetnames==['유류비','변경내역'],book.sheetnames
   rows=list(book['유류비'].iter_rows(values_only=True));assert len(rows)>500
   book.close()
  return elapsed/loops,sig
def train_runs(binary):
 for i in range(16):
  for fixture in [TASK/'before.xlsx',TASK/'valid.xlsx']:
   update(binary,fixture,True,trainenv)
 for i in range(4):
  for fixture in [TASK/'before.xlsx',TASK/'valid.xlsx']:update(binary,fixture,False,trainenv)
 with tempfile.TemporaryDirectory(prefix='fcupdater-macos-training-errors-') as d:
  work=Path(d)
  for args,code in [(['--help'],0),(['--version'],0),(['--bad-option'],1),(['--verify','extra'],1)]:cmd([binary,*args],work,trainenv,expected=code)
  (work/'fuel_cost_chungcheong.xlsx').write_bytes(b'not an xlsx');cmd([binary,'--verify'],work,trainenv,expected=1)
train_runs(train)
report['training']={'verify_updates':32,'skip_updates':8,'cli_cases':4,'malformed_workbook':1,'current_main_workbook_held_out':True,'pipeline':'unchanged default release optimization pipeline; compiler missing-function diagnostics retained and audited','instrumentation':'unchanged actual product CLI, fixed captured response via external DYLD interpose; adapter excluded from product/profdata'}
sysroot=Path(cmd(['rustc','+'+VERSION,'--print','sysroot']).stdout.decode().strip());llvm=sysroot/'lib/rustlib'/TARGET/'bin/llvm-profdata'
profile=ROOT/'pgo'/(TARGET+'.profdata');raws=list(raw.glob('*.profraw'));assert raws
cmd([llvm,'merge','-o',profile,*raws]);shutil.copy2(profile,OUT/profile.name)
(OUT/'profile-functions.txt').write_bytes(cmd([llvm,'show','--all-functions',profile]).stdout)
report['profile']={'sha256':digest(profile),'bytes':profile.stat().st_size,'raw_files':len(raws)};save()
cargo(['build-pgo'],pgo_td)
pgo=pgo_td/TARGET/'release/fcupdater'
cargo(['clippy','--release','--frozen','--bin','fcupdater','--config','.cargo/pgo.toml','--','-D','warnings'],pgo_td)
exes={'release':normal,'pgo':pgo};report['binary']={k:{'bytes':v.stat().st_size,'sha256':digest(v)} for k,v in exes.items()}
for k,v in exes.items():
 otool=cmd(['otool','-L',v]).stdout.decode();assert 'replay' not in otool
 report.setdefault('linkage',{})[k]=otool
with tempfile.TemporaryDirectory(prefix='fcupdater-macos-held-out-errors-') as d:
 work=Path(d)
 for args,code in [(['--help'],0),(['--version'],0),(['--bad-option'],1),(['--verify','extra'],1)]:
  results=[cmd([exe,*args],work,env,expected=code) for exe in exes.values()];assert [(p.returncode,p.stdout,p.stderr) for p in results][0]==[(p.returncode,p.stdout,p.stderr) for p in results][1]
  report['checks'].append('CLI '+str(args))
 for data in [b'not an XLSX',b'PK\x03\x04truncated zip']:
  target=work/'fuel_cost_chungcheong.xlsx';results=[]
  for exe in exes.values():
   target.write_bytes(data);p=cmd([exe,'--verify'],work,env,expected=1);assert target.read_bytes()==data;results.append((p.returncode,p.stdout,p.stderr))
  assert results[0]==results[1];report['checks'].append('Malformed workbook rejection/preservation '+digest_bytes(data))
rng=random.Random(2026101014)
def interval(rows):
 ratios=[r['pgo']/r['release']-1 for r in rows];b=random.Random(2026101015)
 boots=sorted(statistics.median(b.choices(ratios,k=len(ratios))) for _ in range(10000))
 return {'median_seconds':[statistics.median(r[k] for r in rows) for k in exes],'paired_median_change':statistics.median(ratios),'ci95':[boots[250],boots[9749]],'pairs':len(rows),'p95_seconds':[sorted(r[k] for r in rows)[int(.95*len(rows))] for k in exes]}
cases=[('verify-current',ROOT/'fuel_cost_chungcheong.xlsx',True),('verify-prior',TASK/'before.xlsx',True),('skip-current',ROOT/'fuel_cost_chungcheong.xlsx',False),('skip-prior',TASK/'before.xlsx',False)]
for name,fixture,verify in cases:
 sigs=[update(exe,fixture,verify,independent=True)[1] for exe in exes.values()];assert sigs[0]==sigs[1]
 for exe in exes.values():update(exe,fixture,verify,loops=LOOPS)
 data={'raw':[]};report['cases'][name]=data
 for i in range(PAIRS):
  order=list(exes);rng.shuffle(order);values={k:update(exes[k],fixture,verify,loops=LOOPS) for k in order}
  assert values['release'][1]==values['pgo'][1],(name,i,'output mismatch')
  data['raw'].append({'order':order,'release':values['release'][0],'pgo':values['pgo'][0]});save()
  if i%16==0:print(name,'pairs',i+1,'/',PAIRS,flush=True)
 data['summary']=interval(data['raw']);save();print(name,data['summary'],flush=True)
# Per-process native peak RSS, separate from runtime pairs.
report['memory_samples']={}
for label,exe in exes.items():
 with tempfile.TemporaryDirectory(prefix='fcupdater-macos-memory-') as d:
  work=Path(d);shutil.copy2(ROOT/'fuel_cost_chungcheong.xlsx',work/'fuel_cost_chungcheong.xlsx')
  p=cmd(['/usr/bin/time','-l',exe,'--verify'],work,env);text=p.stderr.decode();rss=re.search(r'(\d+)\s+maximum resident set size',text);assert rss,text
  report['memory_samples'][label]={'maximum_resident_set_size_native':int(rss[1]),'raw_time_l':text}
# Clean supported packaging, with no replay dylib or profdata in the archive.
penv=dict(os.environ,CARGO_TARGET_DIR=str(pgo_td/TARGET));penv.pop('DYLD_INSERT_LIBRARIES',None);penv.pop('DYLD_FORCE_FLAT_NAMESPACE',None);penv.pop('LLVM_PROFILE_FILE',None)
suffix='macos-x64' if TARGET.startswith('x86_64') else 'macos-arm64'
cmd(['cargo','+'+VERSION,'run','--frozen','--example','package_artifact','--','fcupdater-'+suffix],env=penv)
package=ROOT/'artifacts'/('fcupdater-'+suffix+'.tar')
with tarfile.open(package) as z:
 assert z.getnames()==['fcupdater-'+suffix];assert z.extractfile(z.getmembers()[0]).read()==pgo.read_bytes()
report['package']={'bytes':package.stat().st_size,'sha256':digest(package)};shutil.copy2(package,OUT/package.name)
report['checks'].append('Clean supported package exact executable equality')
report['primary_improvement']=report['cases']['verify-current']['summary']['paired_median_change']<=-.01 and report['cases']['verify-current']['summary']['ci95'][1]<0
report['non_regression']=all(v['summary']['ci95'][1]<=.03 for v in report['cases'].values())
report['accepted']=report['primary_improvement'] and report['non_regression'];save()
# Bounded live-path checks; latency is not attributed to PGO.
report['live']=[]
for label,exe in exes.items():
 with tempfile.TemporaryDirectory(prefix='fcupdater-macos-live-') as d:
  work=Path(d);shutil.copy2(ROOT/'fuel_cost_chungcheong.xlsx',work/'fuel_cost_chungcheong.xlsx');start=time.perf_counter()
  try:
   p=cmd([exe,'--verify'],work,penv,expected=None,timeout=120)
   result={'variant':label,'exit':p.returncode,'seconds':time.perf_counter()-start,'stdout':p.stdout.decode('utf-8','replace'),'stderr':p.stderr.decode('utf-8','replace')}
   if p.returncode==0:canonical(work/'fuel_cost_chungcheong.xlsx')
  except subprocess.TimeoutExpired:result={'variant':label,'timeout':120}
  report['live'].append(result);save();print('Live',label,result.get('exit'),result.get('timeout'),flush=True)
for path,old in report['linux_windows_profiles'].items():assert digest(ROOT/'pgo'/path)==old
assert report['sources']=={str(p.relative_to(ROOT)):digest(p) for p in sorted((ROOT/'src').rglob('*')) if p.is_file()}
report['complete']=True;save();print('FINAL',TARGET,'accepted',report['accepted'],flush=True)
assert report['accepted'],report['cases']
